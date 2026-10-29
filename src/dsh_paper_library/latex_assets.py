"""Bounded, explicit image intake and inert XLSX values for writing tools."""
from __future__ import annotations

import base64
import binascii
import hashlib
import io
import os
import re
import tempfile
import unicodedata
import zipfile
from pathlib import Path

MAX_BYTES = 8 * 1024 * 1024
MAX_ROWS = 100
MAX_COLUMNS = 30
MAX_CELLS = 2000
MAX_CELL_TEXT = 10000


def _decode(value):
    if not isinstance(value, str) or len(value) > (MAX_BYTES + 2) // 3 * 4:
        raise ValueError("文件超过 8 MiB 上限")
    try:
        body = base64.b64decode(value, validate=True)
    except (ValueError, binascii.Error) as error:
        raise ValueError("文件编码无效") from error
    if not body or len(body) > MAX_BYTES:
        raise ValueError("文件为空或超过 8 MiB 上限")
    return body


def import_image(library, id, name, data_base64):
    from .latex import _project, _resolve_root, _inside, now
    import pymupdf as fitz

    row = _project(library, id)
    root = _resolve_root(row["root"])
    if not isinstance(name, str) or not name or len(name) > 240 or "/" in name or "\\" in name or "\x00" in name:
        raise ValueError("图片文件名无效")
    body = _decode(data_base64)
    if body.startswith(b"\x89PNG\r\n\x1a\n"):
        extension = ".png"
    elif body.startswith(b"\xff\xd8\xff"):
        extension = ".jpg"
    else:
        raise ValueError("仅接收 PNG / JPEG；其他图片请先转换为 PNG")
    try:
        with fitz.open(stream=body) as image:
            if len(image) != 1 or image.is_pdf:
                raise ValueError("不是单张图片")
            profile = image[0].get_image_info()[0]
            width, height = profile["width"], profile["height"]
            if width <= 0 or height <= 0 or width * height > 50_000_000:
                raise ValueError("图片尺寸超过 5000 万像素上限")
    except Exception as error:
        raise ValueError(f"图片无法读取：{error}") from error
    stem = unicodedata.normalize("NFKD", Path(name).stem).encode("ascii", "ignore").decode()
    stem = re.sub(r"[^A-Za-z0-9-]+", "-", stem).strip("-")[:60] or "image"
    digest = hashlib.sha256(body).hexdigest()
    target, relative = _inside(root, f"figures/{stem}-{digest[:12]}{extension}")
    target.parent.mkdir(exist_ok=True)
    # Re-resolve after mkdir as well, including an existing figures symlink.
    target, relative = _inside(root, relative)
    if target.exists():
        if target.is_file() and target.stat().st_size == len(body) and target.read_bytes() == body:
            return {"path": relative, "bytes": len(body), "sha256": digest, "reused": True}
        raise ValueError("同名图片已存在；未覆盖原文件")
    handle, temporary = tempfile.mkstemp(prefix=".paper-image-", dir=target.parent)
    try:
        with os.fdopen(handle, "wb") as stream:
            stream.write(body)
            stream.flush()
            os.fsync(stream.fileno())
        # An exclusive hard link publishes complete bytes without replacing an original.
        os.link(temporary, target)
    except FileExistsError as error:
        raise ValueError("同名图片已存在；未覆盖原文件") from error
    finally:
        os.unlink(temporary)
    library.db.execute("UPDATE latex_projects SET modified=? WHERE id=?", (now(), row["id"]))
    library.db.commit()
    return {"path": relative, "bytes": len(body), "sha256": digest, "reused": False}


def import_table(_library, data_base64, sheet=None):
    """Read one selected worksheet with openpyxl; never evaluate formulas or links."""
    from openpyxl import load_workbook

    body = _decode(data_base64)
    try:
        with zipfile.ZipFile(io.BytesIO(body)) as archive:
            infos = archive.infolist()
            if len(infos) > 1000 or len({item.filename for item in infos}) != len(infos):
                raise ValueError("XLSX 条目过多或重复")
            if any(item.flag_bits & 1 or item.file_size > 8 * 1024 * 1024 for item in infos) or sum(item.file_size for item in infos) > 16 * 1024 * 1024:
                raise ValueError("XLSX 解压内容超过上限或已加密")
            # DTD/entity declarations are unnecessary in an OOXML workbook.
            for item in infos:
                if item.filename.endswith((".xml", ".rels")):
                    xml = archive.read(item)
                    if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                        raise ValueError("XLSX XML 不允许文档类型或实体声明")
    except (zipfile.BadZipFile, ValueError) as error:
        raise ValueError(f"XLSX 无法读取：{error}") from error
    values = formulas = None
    try:
        values = load_workbook(io.BytesIO(body), read_only=True, data_only=True, keep_links=False)
        formulas = load_workbook(io.BytesIO(body), read_only=True, data_only=False, keep_links=False)
        names = values.sheetnames
        if not names or len(names) > 100:
            raise ValueError("XLSX 没有工作表或工作表过多")
        chosen = names[0] if sheet is None else sheet
        if not isinstance(chosen, str) or chosen not in names:
            raise ValueError("指定工作表不存在")
        worksheet, formula_sheet = values[chosen], formulas[chosen]
        if (worksheet.max_row or 0) > MAX_ROWS or (worksheet.max_column or 0) > MAX_COLUMNS:
            raise ValueError("表格超过 100 行 / 30 列上限，请选取较小区域")
        if (worksheet.max_row or 1) * (worksheet.max_column or 1) > MAX_CELLS:
            raise ValueError("表格超过 2000 个单元格上限")
        rows, warnings = [], set()
        for row_index, (data_row, formula_row) in enumerate(zip(worksheet.iter_rows(), formula_sheet.iter_rows()), 1):
            if row_index > MAX_ROWS or len(data_row) > MAX_COLUMNS or row_index * len(data_row) > MAX_CELLS:
                raise ValueError("表格超过 100 行 / 30 列 / 2000 个单元格上限")
            output = []
            for cell, formula_cell in zip(data_row, formula_row):
                if formula_cell.data_type == "f":
                    if cell.value is None:
                        raise ValueError("公式没有缓存结果，请在 Excel 中计算保存后重试或粘贴值")
                    warnings.add("公式仅导入文件中已缓存的值，不重新计算。")
                if cell.data_type == "e":
                    warnings.add("工作表含 Excel 错误值，请插入前核对。")
                value = "" if cell.value is None else str(cell.value)
                if len(value) > MAX_CELL_TEXT:
                    raise ValueError("单元格超过 10000 字符上限")
                output.append(value)
            rows.append(output)
        while rows and not any(rows[-1]):
            rows.pop()
        if not rows or not any(any(row) for row in rows):
            raise ValueError("所选工作表为空")
        return {"sheets": names, "sheet": chosen, "rows": rows, "warnings": sorted(warnings) + ["读取单元格值；字体、颜色、合并单元格及数字显示格式不保留。"]}
    except ValueError:
        raise
    except Exception as error:
        raise ValueError(f"XLSX 无法读取：{error}") from error
    finally:
        if values is not None:
            values.close()
        if formulas is not None:
            formulas.close()
